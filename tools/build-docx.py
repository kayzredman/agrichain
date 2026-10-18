#!/usr/bin/env python3
"""Build the Word copy of the blueprint from docs/blueprint.md."""

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

REPO = Path(__file__).resolve().parent.parent
DOCS = REPO / "docs"
SOURCE = DOCS / "blueprint.md"
TARGET = DOCS / "agriChain Solution Blueprint.docx"
REFERENCE = REPO / "tools" / "reference.docx"
STAGE = Path("/tmp/agrichain-body.docx")
TRIMMED = Path("/tmp/agrichain-blueprint.md")

TITLE = "Solution Blueprint, Technical Architecture & MVP Roadmap"
DATELINE = "17 October 2026"
IMAGE_WIDTH = Inches(5.8333)
BREAK_BEFORE = ("Technical architecture", "MVP roadmap")

INK = RGBColor(0x14, 0x32, 0x28)
ACCENT = RGBColor(0x2D, 0x6A, 0x4F)
DEEP = RGBColor(0x1B, 0x43, 0x32)
BODY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x52, 0x63, 0x6B)


def pandoc():
    """Render the markdown body, dropping its H1 so the cover carries the title."""
    text = SOURCE.read_text(encoding="utf-8")
    text = re.sub(r"\A# .*?\n", "", text, count=1)
    TRIMMED.write_text(text, encoding="utf-8")
    subprocess.run(
        [
            "pandoc",
            str(TRIMMED),
            "--from=markdown",
            "--to=docx",
            f"--reference-doc={REFERENCE}",
            f"--resource-path={DOCS}",
            "--shift-heading-level-by=-1",
            f"--output={STAGE}",
        ],
        check=True,
    )


def line(doc, text, size, color, bold=False, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if text:
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    p.add_run()._r.append(br)
    return p


def contents(doc, headings):
    line(doc, "Contents", 18, INK, bold=True, after=12)
    for level, text in headings:
        if level == 1:
            line(doc, text, 12, INK, bold=True, after=1)
        else:
            line(doc, text, 11, MUTED, after=1)


def build():
    pandoc()
    doc = Document(STAGE)

    headings = [
        (1 if p.style.name == "Heading 1" else 2, p.text)
        for p in doc.paragraphs
        if p.style.name in ("Heading 1", "Heading 2")
    ]

    first = doc.element.body[0]

    # Append the cover and contents, then lift them above the body in reading
    # order. Iterating forward keeps each new paragraph after the previous one.
    cover = [
        line(doc, "", 11, BODY),
        line(doc, "AGRICHAIN", 12, ACCENT, bold=True, after=10),
        line(doc, TITLE, 28, INK, bold=True, after=8),
        line(doc, "", 11, BODY),
        line(doc, "First version  ·  Cashew corridor", 14, DEEP, after=2),
        line(doc, "Côte d'Ivoire to a European buyer", 14, BODY, after=2),
        line(doc, DATELINE, 11, MUTED),
        page_break(doc),
    ]
    contents_start = len(doc.paragraphs)
    contents(doc, headings)
    cover += doc.paragraphs[contents_start:]
    cover.append(page_break(doc))

    for p in cover:
        first.addprevious(p._p)

    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = IMAGE_WIDTH
        shape.height = int(IMAGE_WIDTH * ratio)

    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text in BREAK_BEFORE:
            pPr = p._p.get_or_add_pPr()
            pPr.append(OxmlElement("w:pageBreakBefore"))

    doc.save(TARGET)
    print(f"{TARGET.name}: {len(headings)} headings, {len(doc.inline_shapes)} images")


build()
